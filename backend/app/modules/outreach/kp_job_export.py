"""Excel-выгрузка партии КП целиком: контакты + готовое КП по каждой компании (2026-09-19).

Замечание @user 19.09: отправка из кабинета пока «скоро», а рабочий путь — скачать таблицу
и разослать/обзвонить самим. В отличие от kp_call_list_export («На обзвон» — только
компании без email), здесь все компании партии и все известные каналы: телефон, все email,
сайт, Telegram, WhatsApp, VK — плюс боль, тема и полный текст КП.
"""

from __future__ import annotations

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.maps import Company
from app.modules.outreach.kp_bulk_service import JobItemRow, list_job_items
from app.modules.outreach.kp_call_list_export import _extract_pain_label, _extract_quote
from app.modules.outreach.kp_send_service import collect_company_emails

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="065F46")  # изумруд — цвет бренда

_STATUS_LABELS = {"done": "Готово", "running": "Генерируется", "queued": "В очереди", "failed": "Ошибка"}

_COLUMNS: list[tuple[str, int, bool]] = [
    ("#", 5, False),
    ("Компания", 34, True),
    ("Юрлицо", 30, True),
    ("ИНН", 14, False),
    ("Город", 16, False),
    ("Адрес", 34, True),
    ("Телефон", 20, False),
    ("Email", 32, True),
    ("Сайт", 30, False),
    ("Telegram", 26, True),
    ("WhatsApp", 26, True),
    ("VK", 26, True),
    ("Статус КП", 14, False),
    ("Боль", 28, True),
    ("Цитата из отзыва", 44, True),
    ("Тема КП", 44, True),
    ("Текст КП", 80, True),
]


def _extra_values(extra: dict | None, key: str) -> str:
    """contacts_extra['telegrams'] и т.п. — список строк; склеиваем через перевод строки."""
    if not isinstance(extra, dict):
        return ""
    raw = extra.get(key)
    if isinstance(raw, str):
        raw = [raw]
    if not isinstance(raw, list):
        return ""
    seen: list[str] = []
    for v in raw:
        s = str(v).strip() if v is not None else ""
        if s and s not in seen:
            seen.append(s)
    return "\n".join(seen[:3])


def _row(idx: int, item: JobItemRow, company: tuple | None, emails: list[str]) -> list:
    website, extra = (company[1], company[2]) if company else (None, None)
    draft = item.draft
    return [
        idx,
        item.company_name or f"Компания #{item.company_id or ''}",
        item.company_legal_full or item.company_legal_short or "",
        item.company_inn or "",
        item.company_city or "",
        item.company_address or "",
        item.company_phone or "",
        "\n".join(emails[:3]),
        website or "",
        _extra_values(extra, "telegrams"),
        _extra_values(extra, "whatsapps"),
        _extra_values(extra, "vks"),
        _STATUS_LABELS.get(item.status, item.status),
        _extract_pain_label(draft),
        _extract_quote(draft),
        (draft.subject if draft else "") or "",
        (draft.body if draft else "") or "",
    ]


async def build_job_xlsx(db: AsyncSession, *, user_id: int, job_id: int) -> bytes | None:
    """xlsx всей партии; None — партия не найдена или чужая."""
    result = await list_job_items(db, user_id=user_id, job_id=job_id)
    if result is None:
        return None
    _job, items = result

    company_ids = [int(it.company_id) for it in items if it.company_id is not None]
    companies: dict[int, tuple] = {}
    emails_by_company: dict[int, list[str]] = {}
    if company_ids:
        rows = (
            await db.execute(
                select(Company.id, Company.website, Company.contacts_extra).where(Company.id.in_(company_ids))
            )
        ).all()
        companies = {int(r[0]): tuple(r) for r in rows}
        emails_by_company = await collect_company_emails(db, company_ids)

    wb = Workbook()
    ws = wb.active
    ws.title = "Партия КП"
    for col_idx, (header, width, _wrap) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLUMNS))}1"

    for row_idx, item in enumerate(items, start=2):
        cid = int(item.company_id) if item.company_id is not None else None
        emails = emails_by_company.get(cid, []) if cid is not None else []
        if not emails and item.recipient_email:
            emails = [item.recipient_email]
        values = _row(row_idx - 1, item, companies.get(cid) if cid is not None else None, emails)
        for col_idx, ((_h, _w, wrap), value) in enumerate(zip(_COLUMNS, values), start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(wrap_text=wrap, vertical="top")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_job_filename(job_id: int) -> str:
    date = datetime.utcnow().strftime("%Y-%m-%d")
    return f"kp-partiya-{job_id}_{date}.xlsx"
