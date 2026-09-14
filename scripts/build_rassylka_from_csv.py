"""Сборка рассылки из ЛОКАЛЬНОГО обогащённого CSV (а не из прод-БД).

Зачем отдельный вход: боевые email-адреса живут не в прод-БД (там у компаний
с болью всего ~1900 email), а в локальном мастере, обогащённом краулом сайтов и
яндекс-карточками — `exports/companies_with_pains_carded_enriched.csv` (~7255
email). Прод-скрипт `build_rassylka_exports.py` читает БД и упирается в её
потолок; этот — читает тот же CSV, что дал файлы 6506 email от 04-05.09, но
применяет НОВУЮ логику групп (A/B/C/D1/D2), персональные ссылки на лендинг и
deep-link в бота — как в build_rassylka_exports.py.

Никакой БД/сети: чистый python + openpyxl. MX не проверяем (email уже собран с
живых сайтов; проверка 7000+ доменов офлайн не нужна) — только синтаксис + дедуп.

Вход  (env override):
    RSLK_CSV   exports/companies_with_pains_carded_enriched.csv
Реквизиты оффера (env override, дефолты — боевые):
    PUBLIC_BOT_USERNAME   bolshe_lidov_bot
    PUBLIC_LANDING_URL    https://spinlid.ru/razbor
    PUBLIC_CONTACT_EMAIL  dmitry@spinlid-team.ru
Выход:
    exports/<date>_rassylka/rassylka_email_<date>.xlsx
    exports/<date>_rassylka/rassylka_telegram_<date>.xlsx
Флаги:
    --date YYYY-MM-DD   дата в имени файлов (по умолчанию сегодня)
    --limit N           ограничить число компаний (отладка)
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.modules.outreach.pain_dictionaries import (
    AUTOMATION,
    landing_slug_for_key,
    match_pain_key,
    offer_group_for_key,
)

# Печать кириллицы/эмодзи в cp1251-консоли Windows не должна ронять процесс.
try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # noqa: BLE001
    pass


def _cfg(name: str, default: str) -> str:
    v = os.environ.get(name)
    return v if v else default


BOT = _cfg("PUBLIC_BOT_USERNAME", "bolshe_lidov_bot").lstrip("@").strip()
LANDING = _cfg("PUBLIC_LANDING_URL", "https://spinlid.ru/razbor").strip().rstrip("/")
EMAIL = _cfg("PUBLIC_CONTACT_EMAIL", "dmitry@spinlid-team.ru").strip()

_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]{2,}$")
_PAIN_RE = re.compile(r"^(.*?)\s*\((\d+)\)\s*$")

# Темы письма под каждый pain_key (как в build_rassylka_exports.py).
_SUBJECTS: dict[str, str] = {
    "call_no_answer": "Пропущенные звонки в «{name}»",
    "callback_lost": "Потерянные заявки в «{name}»",
    "chat_no_response": "Сообщения клиентов «{name}» без ответа",
    "schedule_hard": "Запись клиентов в «{name}»",
    "schedule_wait": "Долгое ожидание записи в «{name}»",
    "order_online_hard": "Оформление заказа на сайте «{name}»",
    "order_wait": "«Где мой заказ?» — сроки в «{name}»",
    "queue_wait": "Очереди и ожидание в «{name}»",
}

GIANT_STOPWORDS: tuple[str, ...] = (
    "сбербанк", "сбер", "втб", "альфа-банк", "тинькофф", "т-банк", "росбанк",
    "газпромбанк", "райффайзен", "открытие банк", "совкомбанк", "почта банк",
    "мтс", "билайн", "мегафон", "теле2", "tele2", "ростелеком", "yota", "йота",
    "пятёрочка", "пятерочка", "магнит", "перекрёсток", "перекресток", "дикси",
    "лента", "ашан", "auchan", "метро cash", "окей ", "вкусвилл", "wildberries",
    "вайлдберриз", "озон", "ozon", "яндекс маркет", "яндекс.маркет", "днс ",
    "эльдорадо", "м.видео", "мвидео", "leroy", "леруа", "спортмастер",
    "мфц", "гбуз", "гауз", "фгбу", "фку", "мвд", "администрация", "департамент",
    "министерств", "росреестр", "пенсионный фонд", "налоговая", "фнс ",
    "почта россии", "ржд", "аэрофлот",
)


def _is_giant(name: str) -> bool:
    low = (name or "").lower()
    return any(w in low for w in GIANT_STOPWORDS)


def _clip(s: str | None, n: int) -> str:
    s = (s or "").strip().replace("\n", " ")
    return s if len(s) <= n else s[: n - 1] + "…"


def _tg_type(value: str) -> str:
    v = (value or "").strip().lower()
    handle = v.rsplit("/", 1)[-1].lstrip("@")
    if handle.endswith("bot"):
        return "бот"
    if "/+" in v or "joinchat" in v or "/c/" in v:
        return "канал"
    return "личный"


def _split_multi(raw: str) -> list[str]:
    return [p.strip() for p in re.split(r"[;,]", raw or "") if p.strip()]


def _parse_pains(raw: str) -> list[tuple[str, int]]:
    """'боль1 (45); боль2 (44)' → [('боль1', 45), ('боль2', 44)]."""
    out: list[tuple[str, int]] = []
    for chunk in _split_multi(raw):
        m = _PAIN_RE.match(chunk)
        if m:
            out.append((m.group(1).strip(), int(m.group(2))))
        elif chunk:
            out.append((chunk, 0))
    return out


def _dominant_key(pains: list[tuple[str, int]]) -> str | None:
    keycount: dict[str, int] = {}
    for label, m in pains:
        k = match_pain_key(label)
        if k:
            keycount[k] = keycount.get(k, 0) + (m or 1)
    return max(keycount, key=keycount.get) if keycount else None


# ── Персональные ссылки (как в build_rassylka_exports.py) ─────────────────────


def _landing_link(cid: str, key: str | None) -> str:
    slug = landing_slug_for_key(key)
    url = f"{LANDING}/{slug}" if slug else LANDING
    return f"{url}?c={cid}&utm_source=email"


def _bot_deeplink(cid: str, key: str | None) -> str:
    group = landing_slug_for_key(key) or "general"
    return f"https://t.me/{BOT}?start={group}_{cid}"


def _contact_block_email(cid: str, key: str | None) -> list[str]:
    return [
        "—",
        "Дмитрий",
        f"Разбор по вашей компании: {_landing_link(cid, key)}",
        f"Или сразу в Telegram: {_bot_deeplink(cid, key)}",
        f"Почта: {EMAIL}",
    ]


def build_email_text(name: str, key: str | None, quote: str | None, pain_label: str, cid: str) -> tuple[str, str]:
    subj = (_SUBJECTS.get(key or "") or "Про клиентов, которые не дошли до «{name}»").format(name=name)
    cons = AUTOMATION.consequence.get(key) if key else None
    sol = AUTOMATION.solution.get(key) if key else None
    if quote:
        opener = f"Здравствуйте! Смотрел отзывы о «{name}» — клиент пишет: «{_clip(quote, 280)}»."
    else:
        opener = f"Здравствуйте! Смотрел отзывы о «{name}» и вижу повторяющуюся проблему: {pain_label.lower()}."
    lines = [opener]
    if cons:
        lines.append(f"{cons.capitalize()}.")
    if sol:
        lines.append(f"Обычно помогает вот что: {sol}.")
    lines.append("Могу за 5 минут показать на вашем примере, как это выглядит, — когда удобно ответить?")
    lines.append("")
    lines.extend(_contact_block_email(cid, key))
    return subj, "\n".join(lines)


def build_tg_text(name: str, key: str | None, quote: str | None, pain_label: str) -> str:
    cons = AUTOMATION.consequence.get(key) if key else None
    sol = AUTOMATION.solution.get(key) if key else None
    if quote:
        opener = f"Здравствуйте! Смотрел отзывы о «{name}» — клиент пишет: «{_clip(quote, 200)}»."
    else:
        opener = f"Здравствуйте! Смотрел отзывы о «{name}»: повторяется проблема — {pain_label.lower()}."
    lines = [opener]
    if cons:
        lines.append(f"{cons.capitalize()}.")
    if sol:
        lines.append(f"Обычно помогает так: {sol}.")
    lines.append(f"Могу показать на вашем примере — ответить можно здесь или в Telegram @{BOT}. Когда удобно?")
    return "\n".join(lines)


# ── xlsx ──────────────────────────────────────────────────────────────────────

_HFONT = Font(bold=True, color="FFFFFF")
_HFILL = PatternFill("solid", fgColor="1E3A5F")

_EMAIL_COLS = [
    ("Название", 32, "name"), ("Ниша", 20, "niche"), ("Город", 15, "city"),
    ("Рейтинг", 8, "rating"), ("Отзывов", 8, "reviews"), ("Негатив", 8, "neg"),
    ("Температура", 11, "temp"), ("Группа", 8, "group"),
    ("Боли (с счётчиками)", 42, "pains"), ("Цитата (коротко)", 50, "quote"),
    ("Email", 30, "contact"), ("Тема письма", 34, "subject"),
    ("Текст письма (готовый)", 80, "body"), ("Статус", 11, "_st"),
    ("Дата отправки", 13, "_dt"), ("Ответил", 9, "_rep"), ("Заявка", 9, "_deal"),
    ("Комментарий", 24, "_cmt"),
]
_EMAIL_WRAP = {"Боли (с счётчиками)", "Цитата (коротко)", "Текст письма (готовый)"}

_TG_COLS = [
    ("Название", 32, "name"), ("Ниша", 20, "niche"), ("Город", 15, "city"),
    ("Рейтинг", 8, "rating"), ("Отзывов", 8, "reviews"), ("Негатив", 8, "neg"),
    ("Температура", 11, "temp"), ("Группа", 8, "group"), ("Боли", 42, "pains"),
    ("Цитата (коротко)", 50, "quote"),
    ("Telegram (@username / ссылка)", 28, "contact"),
    ("Тип (личный/канал/бот)", 16, "tg_type"),
    ("Текст сообщения (готовый, без ссылок)", 70, "body"), ("Статус", 11, "_st"),
    ("Дата отправки", 13, "_dt"), ("Ответил", 9, "_rep"), ("Заявка", 9, "_deal"),
    ("Комментарий", 24, "_cmt"),
]
_TG_WRAP = {"Боли", "Цитата (коротко)", "Текст сообщения (готовый, без ссылок)"}


def _write_xlsx(path: Path, title: str, columns: list, rows: list[dict], wrap: set) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Рассылка"
    ws.append([f"{title} · контактов: {len(rows)}"])
    ws.append([])
    hidx = ws.max_row + 1
    ws.append([c[0] for c in columns])
    for i, (_t, w, _k) in enumerate(columns, start=1):
        cell = ws.cell(row=hidx, column=i)
        cell.font = _HFONT
        cell.fill = _HFILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in rows:
        ws.append([r.get(k) for _t, _w, k in columns])
        rr = ws.max_row
        for i, (t, _w, _k) in enumerate(columns, start=1):
            if t in wrap:
                ws.cell(row=rr, column=i).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = ws.cell(row=hidx + 1, column=1)
    wb.save(path)


def _num(v: str):
    v = (v or "").strip()
    if not v:
        return None
    try:
        return float(v) if "." in v else int(v)
    except ValueError:
        return None


def _base(row: dict, pains_txt: str, quote: str, group: str) -> dict:
    return {
        "name": row.get("name"),
        "niche": row.get("niche"),
        "city": row.get("city"),
        "rating": _num(row.get("rating")),
        "reviews": _num(row.get("reviews_count")),
        "neg": _num(row.get("reviews_negative_count")),
        "temp": _num(row.get("lead_temperature")),
        "group": group,
        "pains": pains_txt,
        "quote": _clip(quote, 300),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--date", default="", help="дата в имени файлов (YYYY-MM-DD)")
    parser.add_argument("--limit", type=int, default=0)
    args = parser.parse_args()

    csv.field_size_limit(10_000_000)
    src = Path(_cfg("RSLK_CSV", "exports/companies_with_pains_carded_enriched.csv"))
    stamp = args.date.strip() or date.today().isoformat()
    out_dir = Path("exports") / f"{stamp}_rassylka"
    out_dir.mkdir(parents=True, exist_ok=True)

    with src.open(encoding="utf-8-sig", newline="") as fh:
        rows = list(csv.DictReader(fh))

    # Сорт по температуре ↓ (пусто/None в конец).
    def _tkey(r: dict) -> int:
        t = _num(r.get("lead_temperature"))
        return t if isinstance(t, int) else (int(t) if isinstance(t, float) else -1)

    rows.sort(key=_tkey, reverse=True)
    if args.limit:
        rows = rows[: args.limit]

    stats = {"giants": 0, "no_email": 0, "bad_email": 0, "dup_email": 0, "no_tg": 0, "dup_tg": 0, "no_pain": 0}
    email_rows: list[dict] = []
    tg_rows: list[dict] = []
    seen_email: set[str] = set()
    seen_tg: set[str] = set()

    for row in rows:
        name = (row.get("name") or "").strip()
        cid = (row.get("id") or "").strip()
        if not name or not cid:
            continue
        if _is_giant(name):
            stats["giants"] += 1
            continue
        pains = _parse_pains(row.get("pains") or "")
        if not pains:
            stats["no_pain"] += 1
            continue
        key = _dominant_key(pains)
        group = offer_group_for_key(key)
        pains_txt = "; ".join(f"{lbl} ({m})" for lbl, m in pains)
        quote = (row.get("top_quote") or "").strip()
        first_pain = pains[0][0]

        # ── email (все валидные адреса компании, глобальный дедуп по адресу) ──
        raw_emails = _split_multi(row.get("emails") or "")
        valid_emails = [e for e in raw_emails if _EMAIL_RE.match(e)]
        if not valid_emails:
            stats["bad_email" if raw_emails else "no_email"] += 1
        subj = body = None
        for e in valid_emails:
            el = e.lower()
            if el in seen_email:
                stats["dup_email"] += 1
                continue
            seen_email.add(el)
            if subj is None:
                subj, body = build_email_text(name, key, quote or None, first_pain, cid)
            rec = _base(row, pains_txt, quote, group)
            rec.update({"contact": e, "subject": subj, "body": body})
            email_rows.append(rec)

        # ── telegram (все хэндлы компании, глобальный дедуп) ──
        tgs = _split_multi(row.get("telegram") or "")
        if not tgs:
            stats["no_tg"] += 1
        tg_body = None
        for tg in tgs:
            tl = tg.lower()
            if tl in seen_tg:
                stats["dup_tg"] += 1
                continue
            seen_tg.add(tl)
            if tg_body is None:
                tg_body = build_tg_text(name, key, quote or None, first_pain)
            rec = _base(row, pains_txt, quote, group)
            rec.update({"contact": tg, "tg_type": _tg_type(tg), "body": tg_body})
            tg_rows.append(rec)

    email_path = out_dir / f"rassylka_email_{stamp}.xlsx"
    tg_path = out_dir / f"rassylka_telegram_{stamp}.xlsx"
    _write_xlsx(email_path, "AI-автоматизация приёма клиентов · email", _EMAIL_COLS, email_rows, _EMAIL_WRAP)
    _write_xlsx(tg_path, "AI-автоматизация приёма клиентов · Telegram", _TG_COLS, tg_rows, _TG_WRAP)

    def _grp(rows_: list[dict]) -> str:
        by: dict[str, int] = {}
        for r in rows_:
            by[r["group"]] = by.get(r["group"], 0) + 1
        return ", ".join(f"{g}={by[g]}" for g in sorted(by))

    print(f"[csv-rassylka] источник: {src} · строк CSV: {len(rows)}")
    print(f"[csv-rassylka] email    : {len(email_rows):5d} -> {email_path}")
    print(f"[csv-rassylka] telegram : {len(tg_rows):5d} -> {tg_path}")
    print(f"[csv-rassylka] email по группам   : {_grp(email_rows)}")
    print(f"[csv-rassylka] telegram по группам : {_grp(tg_rows)}")
    print(
        f"[csv-rassylka] отсеяно: гиганты={stats['giants']}, без_боли={stats['no_pain']}, "
        f"без_email={stats['no_email']}, невалид_email={stats['bad_email']}, "
        f"дубль_email={stats['dup_email']}, без_tg={stats['no_tg']}, дубль_tg={stats['dup_tg']}"
    )


if __name__ == "__main__":
    main()
